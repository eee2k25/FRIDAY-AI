"""PowerPoint tools — build, read and edit .pptx decks via python-pptx.

The flagship tool is create_presentation(), which turns a single markdown-lite
outline into an entire deck. That keeps FRIDAY to one tool call for "make me a
deck about X" instead of a slide-by-slide conversation:

    # Quarterly Results          ← each '# ' starts a new slide
    - Revenue up 24%             ← bullets
      - Driven by EMEA           ← indent two spaces for a sub-bullet
    Notes: open with the EMEA story   ← speaker notes for this slide

    # Numbers
    | Region | Q3 |              ← a pipe table becomes a real PowerPoint table
    |---|---|
    | EMEA | 4.1M |

Cross-platform and Office-free. pptx_to_pdf() routes through core.office, so it
uses real PowerPoint on Windows and falls back to LibreOffice elsewhere.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

# Default deck geometry: 16:9 widescreen, in inches.
WIDESCREEN = (13.333, 7.5)
STANDARD = (10.0, 7.5)

LAYOUT_TITLE = 0
LAYOUT_BULLETS = 1
LAYOUT_SECTION = 2
LAYOUT_TITLE_ONLY = 5
LAYOUT_BLANK = 6

LAYOUT_ALIASES = {
    "title": LAYOUT_TITLE,
    "bullets": LAYOUT_BULLETS,
    "content": LAYOUT_BULLETS,
    "section": LAYOUT_SECTION,
    "title_only": LAYOUT_TITLE_ONLY,
    "blank": LAYOUT_BLANK,
}

CHART_TYPES = ("bar", "column", "line", "pie", "doughnut")

MAX_READ_CHARS = 40000


def _p(path: str) -> Path:
    p = Path(path).expanduser()
    if not p.is_absolute():
        p = Path.cwd() / p
    return p


def _pptx():
    try:
        import pptx

        return pptx
    except ImportError as e:
        raise RuntimeError("python-pptx not installed. Run: pip install python-pptx") from e


def _open_deck(path: str):
    """Load an existing .pptx with helpful errors."""
    pptx = _pptx()
    p = _p(path)
    if not p.exists():
        raise FileNotFoundError(f"presentation not found: {p}")
    if p.suffix.lower() != ".pptx":
        raise ValueError(f"not a PowerPoint file: {p.name} (need .pptx; .ppt is not supported)")
    try:
        return pptx.Presentation(str(p)), p
    except Exception as e:  # noqa: BLE001
        raise RuntimeError(f"could not open {p.name}: {type(e).__name__}: {e}") from e


def _resolve_layout(prs, layout) -> int:
    """Accept a layout name ('bullets') or index and return a valid index."""
    if isinstance(layout, str):
        key = layout.strip().lower().replace(" ", "_").replace("-", "_")
        if key in LAYOUT_ALIASES:
            return LAYOUT_ALIASES[key]
        for i, lay in enumerate(prs.slide_layouts):
            if lay.name.strip().lower() == layout.strip().lower():
                return i
        raise ValueError(
            f"unknown layout {layout!r}. Use one of {sorted(LAYOUT_ALIASES)} or a layout index."
        )
    idx = int(layout)
    if not 0 <= idx < len(prs.slide_layouts):
        raise ValueError(f"layout index {idx} out of range (0-{len(prs.slide_layouts) - 1})")
    return idx


def _set_title(slide, title: str) -> None:
    if not title:
        return
    holder = slide.shapes.title
    if holder is not None:
        holder.text = title


def _body_placeholder(slide):
    """The first non-title placeholder that accepts text."""
    for shape in slide.placeholders:
        if shape.placeholder_format.idx != 0 and shape.has_text_frame:
            return shape
    return None


# ------------------------------------------------------------- bullets ---
def _bullet_level(line: str) -> int:
    """Indentation → bullet depth. Two spaces or one tab per level, max 4."""
    stripped = line.lstrip("\t ")
    indent = len(line) - len(stripped)
    tabs = line[:indent].count("\t")
    spaces = indent - tabs
    return min(4, tabs + spaces // 2)


def _fill_bullets(text_frame, lines: list[str]) -> int:
    """Write bullet lines (with levels and **bold**) into a text frame."""
    text_frame.clear()
    written = 0
    for line in lines:
        level = _bullet_level(line)
        text = re.sub(r"^[\-\*\u2022]\s+", "", line.strip())
        if not text:
            continue
        para = text_frame.paragraphs[0] if written == 0 else text_frame.add_paragraph()
        para.level = level
        for seg in re.split(r"(\*\*.+?\*\*)", text):
            if not seg:
                continue
            run = para.add_run()
            if seg.startswith("**") and seg.endswith("**") and len(seg) > 4:
                run.text = seg[2:-2]
                run.font.bold = True
            else:
                run.text = seg
        written += 1
    return written


# -------------------------------------------------------------- tables ---
def _is_table_row(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.endswith("|") and len(s) > 2


def _split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _is_separator_row(line: str) -> bool:
    if not _is_table_row(line):
        return False
    return all(re.fullmatch(r":?-{2,}:?", c) for c in _split_row(line) if c != "")


def _parse_rows(data) -> list[list]:
    if isinstance(data, str):
        try:
            data = json.loads(data)
        except json.JSONDecodeError as e:
            raise ValueError(f'rows must be JSON, e.g. [["Region","Q3"],["EMEA",4.1]]: {e}') from e
    if not isinstance(data, list) or not data:
        raise ValueError("rows must be a non-empty list of rows")
    return [r if isinstance(r, list) else [r] for r in data]


def _add_table_shape(prs, slide, rows: list[list]):
    """Place a table in the slide body area with a bold header row."""
    from pptx.util import Inches, Pt

    width = max(len(r) for r in rows)
    left, top = Inches(0.6), Inches(1.8)
    shape_w = prs.slide_width - Inches(1.2)
    shape_h = Inches(0.4) * len(rows)
    table = slide.shapes.add_table(len(rows), width, left, top, shape_w, shape_h).table
    for ri, row in enumerate(rows):
        for ci in range(width):
            cell = table.cell(ri, ci)
            cell.text = "" if ci >= len(row) else str(row[ci])
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(14)
                    if ri == 0:
                        run.font.bold = True
    return table


# ------------------------------------------------------ outline parsing ---
def parse_outline(outline: str) -> list[dict]:
    """Split a markdown-lite outline into slide dicts.

    Returns [{title, bullets: [...], table: [[...]], notes: str}].
    Content before the first '# ' heading is attached to an untitled slide.
    """
    slides: list[dict] = []
    current: dict | None = None
    notes_mode = False

    def _new(title: str) -> dict:
        return {"title": title, "bullets": [], "table": [], "notes": ""}

    for raw in str(outline).splitlines():
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("# "):
            if current is not None:
                slides.append(current)
            current = _new(stripped[2:].strip())
            notes_mode = False
            continue

        if current is None:
            if not stripped:
                continue
            current = _new("")

        if re.match(r"^notes?\s*:", stripped, re.IGNORECASE):
            current["notes"] = re.sub(r"^notes?\s*:\s*", "", stripped, flags=re.IGNORECASE)
            notes_mode = True
            continue

        if not stripped:
            notes_mode = False
            continue

        if notes_mode:
            current["notes"] += " " + stripped
            continue

        if _is_table_row(line):
            if not _is_separator_row(line):
                current["table"].append(_split_row(line))
            continue

        current["bullets"].append(line)

    if current is not None:
        slides.append(current)
    return slides


def _build_slide(prs, spec: dict):
    """Create one slide from a parsed outline entry."""
    has_table = bool(spec["table"])
    has_bullets = bool(spec["bullets"])
    layout_idx = LAYOUT_TITLE_ONLY if has_table else (LAYOUT_BULLETS if has_bullets else LAYOUT_SECTION)
    slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
    _set_title(slide, spec["title"])

    if has_bullets and not has_table:
        body = _body_placeholder(slide)
        if body is not None:
            _fill_bullets(body.text_frame, spec["bullets"])
    elif has_bullets and has_table:
        # bullets first, then the table underneath
        from pptx.util import Inches

        box = slide.shapes.add_textbox(
            Inches(0.6), Inches(1.4), prs.slide_width - Inches(1.2), Inches(1.2)
        )
        _fill_bullets(box.text_frame, spec["bullets"])

    if has_table:
        _add_table_shape(prs, slide, spec["table"])

    if spec["notes"]:
        slide.notes_slide.notes_text_frame.text = spec["notes"].strip()
    return slide


# ---------------------------------------------------------------- tools ---
def create_presentation(
    path: str, title: str, outline: str = "", subtitle: str = "", widescreen: bool = True
) -> str:
    """Create a .pptx deck: a title slide plus one slide per '# ' in the outline."""
    pptx = _pptx()
    from pptx.util import Inches

    p = _p(path)
    p.parent.mkdir(parents=True, exist_ok=True)

    prs = pptx.Presentation()
    w, h = WIDESCREEN if widescreen else STANDARD
    prs.slide_width, prs.slide_height = Inches(w), Inches(h)

    title_slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_TITLE])
    _set_title(title_slide, title)
    if subtitle:
        holder = _body_placeholder(title_slide)
        if holder is not None:
            holder.text = subtitle

    specs = parse_outline(outline) if outline.strip() else []
    for spec in specs:
        _build_slide(prs, spec)

    prs.save(str(p))
    return f"Created: {p} ({len(prs.slides)} slides, {p.stat().st_size:,} bytes)"


def add_slide(path: str, title: str, content: str = "", layout: str = "bullets") -> str:
    """Append one slide to an existing .pptx. Content is bullets or a pipe table."""
    prs, p = _open_deck(path)
    specs = parse_outline(content) if content.strip() else []
    spec = specs[0] if specs else {"title": "", "bullets": [], "table": [], "notes": ""}
    spec["title"] = title

    if spec["table"] or layout == "bullets":
        _build_slide(prs, spec)
    else:
        idx = _resolve_layout(prs, layout)
        slide = prs.slides.add_slide(prs.slide_layouts[idx])
        _set_title(slide, title)
        if spec["bullets"]:
            body = _body_placeholder(slide)
            if body is not None:
                _fill_bullets(body.text_frame, spec["bullets"])
        if spec["notes"]:
            slide.notes_slide.notes_text_frame.text = spec["notes"]
    prs.save(str(p))
    return f"Added slide {len(prs.slides)} ({title!r}) to {p}"


def add_image_slide(path: str, title: str, image_path: str, caption: str = "") -> str:
    """Append a slide containing a centred, aspect-correct image."""
    from pptx.util import Inches, Pt

    prs, p = _open_deck(path)
    img = _p(image_path)
    if not img.exists():
        raise FileNotFoundError(f"image not found: {img}")

    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_TITLE_ONLY])
    _set_title(slide, title)

    top = Inches(1.6)
    avail_w = prs.slide_width - Inches(1.2)
    avail_h = prs.slide_height - top - Inches(1.0 if caption else 0.5)
    try:
        pic = slide.shapes.add_picture(str(img), Inches(0.6), top, width=avail_w)
    except Exception as e:  # noqa: BLE001
        raise RuntimeError(f"could not insert {img.name}: {type(e).__name__}: {e}") from e
    if pic.height > avail_h:  # too tall — rescale by height and re-centre
        ratio = avail_h / pic.height
        pic.height = int(avail_h)
        pic.width = int(pic.width * ratio)
    pic.left = int((prs.slide_width - pic.width) / 2)

    if caption:
        box = slide.shapes.add_textbox(
            Inches(0.6), pic.top + pic.height + Inches(0.1), avail_w, Inches(0.5)
        )
        para = box.text_frame.paragraphs[0]
        run = para.add_run()
        run.text = caption
        run.font.italic = True
        run.font.size = Pt(14)

    prs.save(str(p))
    return f"Added image slide {len(prs.slides)} ({img.name}) to {p}"


def add_table_slide(path: str, title: str, rows: str) -> str:
    """Append a slide with a table built from JSON rows; first row is the header."""
    prs, p = _open_deck(path)
    parsed = _parse_rows(rows)
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_TITLE_ONLY])
    _set_title(slide, title)
    _add_table_shape(prs, slide, parsed)
    prs.save(str(p))
    return f"Added a {len(parsed)}x{max(len(r) for r in parsed)} table slide to {p}"


def add_chart_slide(
    path: str, title: str, categories: str, series: str, chart_type: str = "bar"
) -> str:
    """Append a native, editable chart slide (bar/column/line/pie/doughnut)."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    kind = str(chart_type).strip().lower()
    if kind not in CHART_TYPES:
        raise ValueError(f"chart_type must be one of {list(CHART_TYPES)}, got {chart_type!r}")

    cats = _parse_rows(categories)
    cats = [str(c[0]) if isinstance(c, list) else str(c) for c in cats]

    if isinstance(series, str):
        try:
            series_obj = json.loads(series)
        except json.JSONDecodeError as e:
            raise ValueError(f'series must be JSON, e.g. {{"2024": [1,2,3]}}: {e}') from e
    else:
        series_obj = series
    if not isinstance(series_obj, dict) or not series_obj:
        raise ValueError('series must be a non-empty JSON object, e.g. {"Revenue": [1,2,3]}')

    chart_data = CategoryChartData()
    chart_data.categories = cats
    for name, values in series_obj.items():
        if not isinstance(values, list):
            raise ValueError(f"series {name!r} must map to a list of numbers")
        if len(values) != len(cats):
            raise ValueError(
                f"series {name!r} has {len(values)} values but there are {len(cats)} categories"
            )
        chart_data.add_series(str(name), [float(v) for v in values])

    enum_name = {
        "bar": "BAR_CLUSTERED",
        "column": "COLUMN_CLUSTERED",
        "line": "LINE_MARKERS",
        "pie": "PIE",
        "doughnut": "DOUGHNUT",
    }[kind]

    prs, p = _open_deck(path)
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_TITLE_ONLY])
    _set_title(slide, title)
    slide.shapes.add_chart(
        getattr(XL_CHART_TYPE, enum_name),
        Inches(0.8),
        Inches(1.7),
        prs.slide_width - Inches(1.6),
        prs.slide_height - Inches(2.4),
        chart_data,
    )
    prs.save(str(p))
    return f"Added a {kind} chart slide ({len(series_obj)} series, {len(cats)} categories) to {p}"


def read_presentation(path: str, include_notes: bool = True) -> str:
    """Extract every slide's title, bullets, tables and speaker notes as text."""
    prs, p = _open_deck(path)
    out: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
        out.append(f"\n--- Slide {i}{': ' + title if title else ''} ---")
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(r.text for r in para.runs).strip() or para.text.strip()
                    if text:
                        out.append(f"{'  ' * para.level}- {text}")
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    out.append("| " + " | ".join(c.text.strip() for c in row.cells) + " |")
            if getattr(shape, "has_chart", False):
                chart = shape.chart
                names = ", ".join(s.name or "(unnamed)" for s in chart.plots[0].series)
                out.append(f"[Chart: {chart.chart_type}, series: {names}]")
        if include_notes and slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.append(f"Notes: {notes}")
    text = "\n".join(out).strip() or "(empty presentation)"
    if len(text) > MAX_READ_CHARS:
        text = text[:MAX_READ_CHARS] + f"\n... [truncated, {len(text)} chars total]"
    return text


def set_speaker_notes(path: str, slide_number: int, notes: str) -> str:
    """Set the speaker notes on one slide (1-based)."""
    prs, p = _open_deck(path)
    slides = list(prs.slides)
    n = int(slide_number)
    if not 1 <= n <= len(slides):
        raise ValueError(f"slide {n} does not exist — the deck has {len(slides)} slide(s)")
    slides[n - 1].notes_slide.notes_text_frame.text = str(notes)
    prs.save(str(p))
    return f"Set speaker notes on slide {n} of {p}"


def get_presentation_info(path: str) -> str:
    """Summarise a .pptx: slide count, size, per-slide titles and contents."""
    prs, p = _open_deck(path)
    slides = list(prs.slides)
    from pptx.util import Emu

    images = tables = charts = noted = words = 0
    titles: list[str] = []
    for i, slide in enumerate(slides, 1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
        marks = []
        for shape in slide.shapes:
            if shape.shape_type is not None and "PICTURE" in str(shape.shape_type):
                images += 1
                marks.append("image")
            if getattr(shape, "has_table", False):
                tables += 1
                marks.append("table")
            if getattr(shape, "has_chart", False):
                charts += 1
                marks.append("chart")
            if shape.has_text_frame:
                words += len(shape.text_frame.text.split())
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            noted += 1
            marks.append("notes")
        suffix = f"  [{', '.join(sorted(set(marks)))}]" if marks else ""
        titles.append(f"  {i}. {title or '(untitled)'}{suffix}")

    return "\n".join(
        [
            f"File:        {p}",
            f"Size:        {p.stat().st_size:,} bytes",
            f"Slides:      {len(slides)}",
            f"Dimensions:  {Emu(prs.slide_width).inches:.2f} x {Emu(prs.slide_height).inches:.2f} in"
            f" ({'16:9 widescreen' if Emu(prs.slide_width).inches > 11 else '4:3 standard'})",
            f"Words:       {words:,}",
            f"Images:      {images}   Tables: {tables}   Charts: {charts}",
            f"With notes:  {noted}/{len(slides)}",
            "\nSlides:",
            *titles[:80],
        ]
    )


def pptx_to_pdf(path: str, output_path: str = "") -> str:
    """Convert a .pptx to PDF using PowerPoint if available, else LibreOffice."""
    from core.office import export_pdf

    p = _p(path)
    if not p.exists():
        raise FileNotFoundError(f"presentation not found: {p}")
    pdf = export_pdf("powerpoint", p, output_path or None)
    return f"Exported PDF: {pdf} ({Path(pdf).stat().st_size:,} bytes)"


_DECLARATIONS: list[dict] = [
    {
        "name": "create_presentation",
        "description": (
            "Create a complete PowerPoint .pptx deck in one call. The outline is markdown-lite: "
            "each '# ' line starts a new slide, '- ' lines are bullets (indent two spaces for "
            "sub-bullets), '**bold**' works, a 'Notes: ...' line adds speaker notes, and a "
            "| a | b | pipe table becomes a real PowerPoint table. Prefer this over adding "
            "slides one at a time."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "Output .pptx path"},
                "title": {"type": "string", "description": "Title-slide heading"},
                "outline": {"type": "string", "description": "Markdown-lite slide outline"},
                "subtitle": {"type": "string", "description": "Title-slide subtitle"},
                "widescreen": {"type": "boolean", "description": "16:9 (default) or 4:3"},
            },
            "required": ["path", "title"],
        },
    },
    {
        "name": "add_slide",
        "description": "Append one slide to an existing .pptx, with bullets or a pipe table as content.",
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string"},
                "title": {"type": "string"},
                "content": {"type": "string", "description": "Bullets and/or a pipe table"},
                "layout": {
                    "type": "string",
                    "description": "title, bullets (default), section, title_only or blank",
                },
            },
            "required": ["path", "title"],
        },
    },
    {
        "name": "add_image_slide",
        "description": "Append a slide with an image scaled to fit the slide, plus an optional caption.",
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string"},
                "title": {"type": "string"},
                "image_path": {"type": "string"},
                "caption": {"type": "string"},
            },
            "required": ["path", "title", "image_path"],
        },
    },
    {
        "name": "add_table_slide",
        "description": "Append a slide containing a table built from JSON rows; the first row is a bold header.",
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string"},
                "title": {"type": "string"},
                "rows": {
                    "type": "string",
                    "description": 'JSON list of rows, e.g. [["Region","Q3"],["EMEA",4.1]]',
                },
            },
            "required": ["path", "title", "rows"],
        },
    },
    {
        "name": "add_chart_slide",
        "description": "Append a slide with a native, editable chart (bar, column, line, pie or doughnut).",
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string"},
                "title": {"type": "string"},
                "categories": {
                    "type": "string",
                    "description": 'JSON list of category labels, e.g. ["Q1","Q2","Q3"]',
                },
                "series": {
                    "type": "string",
                    "description": 'JSON object of series name → values, e.g. {"Revenue":[1,2,3]}. Each list must match the number of categories.',
                },
                "chart_type": {
                    "type": "string",
                    "description": "bar (default), column, line, pie or doughnut",
                },
            },
            "required": ["path", "title", "categories", "series"],
        },
    },
    {
        "name": "read_presentation",
        "description": "Extract every slide's title, bullets, tables, charts and speaker notes as text.",
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string"},
                "include_notes": {"type": "boolean", "description": "Default true"},
            },
            "required": ["path"],
        },
    },
    {
        "name": "set_speaker_notes",
        "description": "Set the speaker notes on one slide of a .pptx (slide numbers are 1-based).",
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string"},
                "slide_number": {"type": "integer"},
                "notes": {"type": "string"},
            },
            "required": ["path", "slide_number", "notes"],
        },
    },
    {
        "name": "get_presentation_info",
        "description": "Summarise a .pptx: slide count, dimensions, word count, and a per-slide list showing which have images, tables, charts or notes.",
        "parameters": {
            "type": "object",
            "properties": {"path": {"type": "string"}},
            "required": ["path"],
        },
    },
    {
        "name": "pptx_to_pdf",
        "description": "Convert a .pptx to PDF using Microsoft PowerPoint when available, otherwise LibreOffice.",
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string"},
                "output_path": {"type": "string", "description": "Optional output .pdf path"},
            },
            "required": ["path"],
        },
    },
]


def register_tools(registry) -> None:
    for d in _DECLARATIONS:
        registry.register_tool(d["name"], globals()[d["name"]], d)
