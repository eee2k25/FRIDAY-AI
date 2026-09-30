"""Document tools — Excel (.xlsx) via openpyxl.

Word (.docx) support moved to tools/word_tools.py in v1.3.0 when it grew
tables, find/replace, images and PDF export. The Word functions are
re-exported here so existing imports keep working.
"""
from __future__ import annotations

import json
from pathlib import Path

from tools.word_tools import (  # noqa: F401 — backwards-compatible re-exports
    append_to_word_doc,
    create_project_report,
    create_word_doc,
    read_word_doc,
)


def _p(path: str) -> Path:
    p = Path(path).expanduser()
    if not p.is_absolute():
        p = Path.cwd() / p
    return p


def _xlsx():
    try:
        import openpyxl

        return openpyxl
    except ImportError as e:
        raise RuntimeError("openpyxl not installed. Run: pip install openpyxl") from e


# ---------------------------------------------------------------- excel ---
def create_excel(path: str, data: str, sheet_name: str = "Sheet1") -> str:
    """Create an .xlsx. data is a JSON list of rows: [["a","b"],[1,2]] — first
    row becomes a bold header."""
    openpyxl = _xlsx()
    try:
        rows = json.loads(data)
    except json.JSONDecodeError as e:
        raise ValueError(f"data must be a JSON list of rows, e.g. [[\"a\",\"b\"],[1,2]]: {e}") from e
    if not isinstance(rows, list):
        raise ValueError("data must be a list of rows")
    p = _p(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row if isinstance(row, list) else [row])
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True)
    ws.freeze_panes = "A2"
    wb.save(str(p))
    return f"Created: {p} ({len(rows)} rows)"


def read_excel(path: str, sheet_name: str | None = None) -> str:
    """Read an .xlsx sheet as a pipe-separated table."""
    openpyxl = _xlsx()
    p = _p(path)
    if not p.exists():
        raise FileNotFoundError(f"workbook not found: {p}")
    wb = openpyxl.load_workbook(str(p), data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
    lines = []
    for row in ws.iter_rows(values_only=True):
        cells = ["" if v is None else str(v) for v in row]
        if any(c.strip() for c in cells):
            lines.append(" | ".join(cells))
    if not lines:
        return f"Sheet '{ws.title}' is empty."
    return f"Sheet '{ws.title}' ({len(lines)} rows):\n" + "\n".join(lines[:500])


def append_excel_row(path: str, sheet_name: str, row_data: str) -> str:
    """Append one row (JSON list) to an existing sheet."""
    openpyxl = _xlsx()
    try:
        row = json.loads(row_data)
    except json.JSONDecodeError as e:
        raise ValueError(f"row_data must be a JSON list, e.g. [\"a\",1]: {e}") from e
    if not isinstance(row, list):
        raise ValueError("row_data must be a JSON list")
    p = _p(path)
    if not p.exists():
        raise FileNotFoundError(f"workbook not found: {p}")
    wb = openpyxl.load_workbook(str(p))
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
    ws.append(row)
    wb.save(str(p))
    return f"Appended row to '{ws.title}' in {p}"


_DECLARATIONS: list[dict] = [
    {
        "name": "create_excel",
        "description": "Create an Excel .xlsx from a JSON list of rows. First row becomes a bold header.",
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string"},
                "data": {"type": "string", "description": 'JSON list of rows, e.g. [["Name","Qty"],["Widget",3]]'},
                "sheet_name": {"type": "string"},
            },
            "required": ["path", "data"],
        },
    },
    {
        "name": "read_excel",
        "description": "Read an .xlsx sheet as a pipe-separated table.",
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string"},
                "sheet_name": {"type": "string", "description": "optional sheet name; defaults to the active sheet"},
            },
            "required": ["path"],
        },
    },
    {
        "name": "append_excel_row",
        "description": "Append one row (JSON list) to a sheet in an existing .xlsx.",
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string"},
                "sheet_name": {"type": "string"},
                "row_data": {"type": "string", "description": "JSON list, e.g. [\"Widget\",3,9.99]"},
            },
            "required": ["path", "sheet_name", "row_data"],
        },
    },
]


def register_tools(registry) -> None:
    for d in _DECLARATIONS:
        registry.register_tool(d["name"], globals()[d["name"]], d)
