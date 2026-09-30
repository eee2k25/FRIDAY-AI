"""Word tools — create, read, edit and convert .docx via python-docx.

Markdown-lite accepted anywhere content is taken:
    # / ## / ###     headings
    - / *            bullets
    1.               numbered list
    **bold**         bold runs
    > quote          block quote
    ---              horizontal rule (page-width separator)
    | a | b |        real Word tables (with a |---|---| separator row)

Everything here is pure-Python and cross-platform. The two Windows-flavoured
extras (word_to_pdf, open_in_office) route through core.office, which falls
back to LibreOffice and otherwise reports a clear, actionable error.
"""
from __future__ import annotations

import datetime
import json
import re
from pathlib import Path

MAX_READ_CHARS = 40000


def _p(path: str) -> Path:
    p = Path(path).expanduser()
    if not p.is_absolute():
        p = Path.cwd() / p
    return p


def _docx():
    try:
        import docx

        return docx
    except ImportError as e:
        raise RuntimeError("python-docx not installed. Run: pip install python-docx") from e


def _open_doc(path: str):
    """Load an existing .docx, with a helpful error if it is missing/not a docx."""
    docx = _docx()
    p = _p(path)
    if not p.exists():
        raise FileNotFoundError(f"document not found: {p}")
    if p.suffix.lower() != ".docx":
        raise ValueError(f"not a Word document: {p.name} (need .docx; .doc is not supported)")
    try:
        return docx.Document(str(p)), p
    except Exception as e:  # noqa: BLE001
        raise RuntimeError(f"could not open {p.name}: {type(e).__name__}: {e}") from e


# ------------------------------------------------------ markdown-lite ---
def _is_table_row(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.endswith("|") and len(s) > 2


def _split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _is_separator_row(line: str) -> bool:
    if not _is_table_row(line):
        return False
    return all(re.fullmatch(r":?-{2,}:?", c) for c in _split_row(line) if c != "")


def _add_inline(paragraph, text: str, style: str | None = None):
    if style:
        try:
            paragraph.style = style
        except KeyError:  # style not in this template
            pass
    for seg in re.split(r"(\*\*.+?\*\*)", text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**") and len(seg) > 4:
            paragraph.add_run(seg[2:-2]).bold = True
        else:
            paragraph.add_run(seg)
    return paragraph


def _add_table(doc, block: list[str]) -> None:
    """Render a markdown pipe-table block as a real Word table."""
    rows = [_split_row(r) for r in block if not _is_separator_row(r)]
    if not rows:
        return
    width = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    for ri, row in enumerate(rows):
        for ci in range(width):
            cell = table.cell(ri, ci)
            cell.text = ""
            para = cell.paragraphs[0]
            _add_inline(para, row[ci] if ci < len(row) else "")
            if ri == 0:
                for run in para.runs:
                    run.bold = True


def _add_markdown_line(doc, line: str) -> None:
    stripped = line.strip()
    if not stripped:
        doc.add_paragraph("")
    elif stripped in ("---", "***", "___"):
        _add_inline(doc.add_paragraph(), "_" * 60)
    elif stripped.startswith("### "):
        doc.add_heading(stripped[4:], level=3)
    elif stripped.startswith("## "):
        doc.add_heading(stripped[3:], level=2)
    elif stripped.startswith("# "):
        doc.add_heading(stripped[2:], level=1)
    elif stripped.startswith("> "):
        _add_inline(doc.add_paragraph(), stripped[2:], style="Quote")
    elif stripped.startswith(("- ", "* ")):
        _add_inline(doc.add_paragraph(), stripped[2:], style="List Bullet")
    elif re.match(r"^\d+\.\s", stripped):
        _add_inline(doc.add_paragraph(), re.sub(r"^\d+\.\s", "", stripped), style="List Number")
    else:
        _add_inline(doc.add_paragraph(), line)


def render_markdown(doc, content: str) -> None:
    """Render markdown-lite into a python-docx document, tables included."""
    lines = str(content).splitlines()
    i = 0
    while i < len(lines):
        if (
            _is_table_row(lines[i])
            and i + 1 < len(lines)
            and _is_separator_row(lines[i + 1])
        ):
            block = []
            while i < len(lines) and _is_table_row(lines[i]):
                block.append(lines[i])
                i += 1
            _add_table(doc, block)
            continue
        _add_markdown_line(doc, lines[i])
        i += 1


def _parse_rows(data) -> list[list]:
    """Accept a JSON string or a real list; return a list of row lists."""
    if isinstance(data, str):
        try:
            data = json.loads(data)
        except json.JSONDecodeError as e:
            raise ValueError(f'rows must be JSON, e.g. [["Name","Qty"],["Widget",3]]: {e}') from e
    if not isinstance(data, list) or not data:
        raise ValueError("rows must be a non-empty list of rows")
    return [r if isinstance(r, list) else [r] for r in data]


def _parse_sections(sections) -> dict:
    if isinstance(sections, str):
        try:
            sections = json.loads(sections)
        except json.JSONDecodeError as e:
            raise ValueError(f"sections must be a JSON object: {e}") from e
    if not isinstance(sections, dict) or not sections:
        raise ValueError("sections must be a non-empty {name: content} mapping")
    return sections


# --------------------------------------------------------------- core ---
def create_word_doc(path: str, title: str, content: str) -> str:
    """Create a .docx from a title + markdown-lite content."""
    docx = _docx()
    p = _p(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    doc = docx.Document()
    if title:
        doc.add_heading(title, level=0)
    render_markdown(doc, content)
    doc.save(str(p))
    return f"Created: {p} ({p.stat().st_size} bytes)"


def read_word_doc(path: str) -> str:
    """Extract all text from a .docx, preserving heading/bullet/table structure."""
    doc, _ = _open_doc(path)
    lines: list[str] = []
    for para in doc.paragraphs:
        style = (para.style.name or "") if para.style else ""
        text = para.text.strip()
        if not text:
            continue
        if "Heading 1" in style or "Title" in style:
            lines.append(f"# {text}")
        elif "Heading 2" in style:
            lines.append(f"## {text}")
        elif "Heading 3" in style:
            lines.append(f"### {text}")
        elif "List" in style:
            lines.append(f"- {text}")
        elif "Quote" in style:
            lines.append(f"> {text}")
        else:
            lines.append(text)
    for ti, table in enumerate(doc.tables, 1):
        lines.append(f"\n[Table {ti}]")
        for row in table.rows:
            lines.append("| " + " | ".join(c.text.strip() for c in row.cells) + " |")
    out = "\n".join(lines) if lines else "(empty document)"
    if len(out) > MAX_READ_CHARS:
        out = out[:MAX_READ_CHARS] + f"\n... [truncated, {len(out)} chars total]"
    return out


def append_to_word_doc(path: str, content: str) -> str:
    """Append markdown-lite content to an existing .docx."""
    doc, p = _open_doc(path)
    render_markdown(doc, content)
    doc.save(str(p))
    return f"Appended to {p} ({p.stat().st_size} bytes)"


def create_project_report(path: str, title: str, sections: dict) -> str:
    """Create a formatted project report: title, date, TOC, sectioned body,
    page breaks between sections."""
    docx = _docx()
    secs = _parse_sections(sections)
    p = _p(path)
    p.parent.mkdir(parents=True, exist_ok=True)

    doc = docx.Document()
    doc.add_heading(title, level=0)
    doc.add_paragraph(f"Generated by FRIDAY — {datetime.datetime.now():%d %B %Y, %H:%M}")
    doc.add_heading("Table of Contents", level=1)
    for i, name in enumerate(secs.keys(), 1):
        _add_inline(doc.add_paragraph(), f"{i}. {name}", style="List Number")
    doc.add_page_break()

    for i, (name, body) in enumerate(secs.items()):
        doc.add_heading(f"{i + 1}. {name}", level=1)
        render_markdown(doc, str(body))
        if i < len(secs) - 1:
            doc.add_page_break()

    doc.save(str(p))
    return f"Report created: {p} ({len(secs)} sections, {p.stat().st_size} bytes)"


# ------------------------------------------------------------ editing ---
def add_table_to_word(path: str, rows: str, heading: str = "", style: str = "Table Grid") -> str:
    """Append a table to an existing .docx. `rows` is JSON; first row = header."""
    doc, p = _open_doc(path)
    parsed = _parse_rows(rows)
    if heading:
        doc.add_heading(heading, level=2)
    width = max(len(r) for r in parsed)
    table = doc.add_table(rows=len(parsed), cols=width)
    try:
        table.style = style
    except KeyError:
        table.style = "Table Grid"
    for ri, row in enumerate(parsed):
        for ci in range(width):
            cell = table.cell(ri, ci)
            cell.text = ""
            para = cell.paragraphs[0]
            _add_inline(para, "" if ci >= len(row) else str(row[ci]))
            if ri == 0:
                for run in para.runs:
                    run.bold = True
    doc.save(str(p))
    return f"Added a {len(parsed)}x{width} table to {p}"


def add_image_to_word(path: str, image_path: str, width_inches: float = 6.0, caption: str = "") -> str:
    """Insert an image into an existing .docx, optionally with a caption."""
    from docx.shared import Inches

    doc, p = _open_doc(path)
    img = _p(image_path)
    if not img.exists():
        raise FileNotFoundError(f"image not found: {img}")
    try:
        doc.add_picture(str(img), width=Inches(float(width_inches)))
    except Exception as e:  # noqa: BLE001 — unsupported formats raise many types
        raise RuntimeError(f"could not insert {img.name}: {type(e).__name__}: {e}") from e
    if caption:
        para = doc.add_paragraph()
        run = para.add_run(caption)
        run.italic = True
    doc.save(str(p))
    return f"Inserted {img.name} into {p}" + (f" with caption {caption!r}" if caption else "")


def _replace_in_paragraph(paragraph, find: str, replace: str, match_case: bool) -> int:
    """Replace inside one paragraph. Run-level first (keeps formatting),
    paragraph-level as a fallback when the text spans several runs."""
    hits = 0
    flags = 0 if match_case else re.IGNORECASE
    pattern = re.compile(re.escape(find), flags)

    for run in paragraph.runs:
        found = len(pattern.findall(run.text))
        if found:
            run.text = pattern.sub(replace, run.text)
            hits += found
    if hits:
        return hits

    full = paragraph.text
    found = len(pattern.findall(full))
    if not found:
        return 0
    new_text = pattern.sub(replace, full)
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)
    return found


def word_find_replace(path: str, find: str, replace: str, match_case: bool = True) -> str:
    """Find and replace text throughout a .docx — body, tables, headers and footers."""
    if not find:
        raise ValueError("`find` must not be empty")
    doc, p = _open_doc(path)
    total = 0
    for para in doc.paragraphs:
        total += _replace_in_paragraph(para, find, replace, match_case)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    total += _replace_in_paragraph(para, find, replace, match_case)
    for section in doc.sections:
        for container in (section.header, section.footer):
            for para in container.paragraphs:
                total += _replace_in_paragraph(para, find, replace, match_case)
    if not total:
        return f"No occurrences of {find!r} found in {p.name} — nothing changed."
    doc.save(str(p))
    return f"Replaced {total} occurrence(s) of {find!r} with {replace!r} in {p}"


def get_word_doc_info(path: str) -> str:
    """Summarise a .docx: word/paragraph/table/image counts and its heading outline."""
    doc, p = _open_doc(path)
    paragraphs = [para for para in doc.paragraphs if para.text.strip()]
    words = sum(len(para.text.split()) for para in paragraphs)
    outline = []
    for para in doc.paragraphs:
        style = (para.style.name or "") if para.style else ""
        text = para.text.strip()
        if text and ("Heading" in style or "Title" in style):
            depth = 0 if "Title" in style else int(re.sub(r"\D", "", style) or 1)
            outline.append(f"{'  ' * depth}• {text}")
    core = doc.core_properties
    lines = [
        f"File:       {p}",
        f"Size:       {p.stat().st_size:,} bytes",
        f"Paragraphs: {len(paragraphs)}",
        f"Words:      {words:,}",
        f"Tables:     {len(doc.tables)}",
        f"Images:     {len(doc.inline_shapes)}",
        f"Sections:   {len(doc.sections)}",
        f"Author:     {core.author or '(unset)'}",
        f"Modified:   {core.modified or '(unknown)'}",
    ]
    if outline:
        lines.append("\nOutline:")
        lines.extend(outline[:60])
    return "\n".join(lines)


# ---------------------------------------------------- Windows/COM side ---
def word_to_pdf(path: str, output_path: str = "") -> str:
    """Convert a .docx to PDF using Microsoft Word if available, else LibreOffice."""
    from core.office import export_pdf

    p = _p(path)
    if not p.exists():
        raise FileNotFoundError(f"document not found: {p}")
    pdf = export_pdf("word", p, output_path or None)
    size = Path(pdf).stat().st_size
    return f"Exported PDF: {pdf} ({size:,} bytes)"


def open_in_office(path: str) -> str:
    """Open a document in its Office application (Word/Excel/PowerPoint)."""
    from core.office import open_file

    return open_file(_p(path))


def office_status() -> str:
    """Report what Office integration can do on this machine (COM, LibreOffice)."""
    from core.office import availability

    info = availability()
    return (
        f"Platform:        {info['platform']}\n"
        f"Live COM apps:   {'yes' if info['com_automation'] else 'no'}\n"
        f"LibreOffice:     {'yes' if info['libreoffice'] else 'no'}\n"
        f"{info['note']}"
    )


_DECLARATIONS: list[dict] = [
    {
        "name": "create_word_doc",
        "description": (
            "Create a Word .docx from a title and markdown-lite content. Supports # headings, "
            "- bullets, 1. numbering, **bold**, > quotes, and | a | b | pipe tables (with a "
            "|---|---| separator row) which become real Word tables."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "Output .docx path"},
                "title": {"type": "string"},
                "content": {"type": "string", "description": "Markdown-lite body text"},
            },
            "required": ["path", "title", "content"],
        },
    },
    {
        "name": "read_word_doc",
        "description": "Extract all text from a .docx preserving heading/bullet/table structure.",
        "parameters": {
            "type": "object",
            "properties": {"path": {"type": "string"}},
            "required": ["path"],
        },
    },
    {
        "name": "append_to_word_doc",
        "description": "Append markdown-lite content (headings, bullets, tables) to an existing .docx.",
        "parameters": {
            "type": "object",
            "properties": {"path": {"type": "string"}, "content": {"type": "string"}},
            "required": ["path", "content"],
        },
    },
    {
        "name": "create_project_report",
        "description": "Create a full project report .docx: title, date, table of contents, one section per entry, page breaks between sections.",
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "Output .docx path"},
                "title": {"type": "string"},
                "sections": {
                    "type": "object",
                    "description": "Ordered mapping of section name → markdown-lite body text.",
                },
            },
            "required": ["path", "title", "sections"],
        },
    },
    {
        "name": "add_table_to_word",
        "description": "Append a real table to an existing .docx. Rows are JSON; the first row becomes a bold header.",
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string"},
                "rows": {
                    "type": "string",
                    "description": 'JSON list of rows, e.g. [["Item","Qty"],["Widget",3]]',
                },
                "heading": {"type": "string", "description": "Optional heading above the table"},
                "style": {"type": "string", "description": "Word table style (default 'Table Grid')"},
            },
            "required": ["path", "rows"],
        },
    },
    {
        "name": "add_image_to_word",
        "description": "Insert an image (png/jpg/gif) into an existing .docx, optionally with an italic caption.",
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "The .docx to insert into"},
                "image_path": {"type": "string"},
                "width_inches": {"type": "number", "description": "Display width, default 6.0"},
                "caption": {"type": "string"},
            },
            "required": ["path", "image_path"],
        },
    },
    {
        "name": "word_find_replace",
        "description": "Find and replace text throughout a .docx — body, tables, headers and footers. Reports how many occurrences changed.",
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string"},
                "find": {"type": "string"},
                "replace": {"type": "string"},
                "match_case": {"type": "boolean", "description": "Default true"},
            },
            "required": ["path", "find", "replace"],
        },
    },
    {
        "name": "get_word_doc_info",
        "description": "Summarise a .docx: word/paragraph/table/image counts, author, and the heading outline. Use before editing a document you have not read.",
        "parameters": {
            "type": "object",
            "properties": {"path": {"type": "string"}},
            "required": ["path"],
        },
    },
    {
        "name": "word_to_pdf",
        "description": "Convert a .docx to PDF using Microsoft Word when available, otherwise LibreOffice.",
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string"},
                "output_path": {"type": "string", "description": "Optional output .pdf path"},
            },
            "required": ["path"],
        },
    },
    {
        "name": "open_in_office",
        "description": "Open a document in its Office application so the Boss can see it (Windows/macOS/Linux desktop).",
        "parameters": {
            "type": "object",
            "properties": {"path": {"type": "string"}},
            "required": ["path"],
        },
    },
    {
        "name": "office_status",
        "description": "Report which Office integrations are available here: live COM app automation and/or LibreOffice for PDF export.",
        "parameters": {"type": "object", "properties": {}},
    },
]


def register_tools(registry) -> None:
    for d in _DECLARATIONS:
        registry.register_tool(d["name"], globals()[d["name"]], d)
